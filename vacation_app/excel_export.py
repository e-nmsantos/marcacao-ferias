from __future__ import annotations

from calendar import monthrange
from datetime import date, timedelta
from io import BytesIO

from vacation_app.calendar_utils import month_name_pt

PT_WEEKDAYS = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]


def vacations_for_year(vacations: list[dict], year: int, team: str) -> list[dict]:
    approved = []
    for vacation in vacations:
        if vacation.get("status") != "approved":
            continue
        if team != "Todas" and vacation.get("team", "") != team:
            continue
        if vacation["start_date"].year != year and vacation["end_date"].year != year:
            continue
        approved.append(vacation)
    return approved


def employee_names_for_map(staff: list[dict], vacations: list[dict], team: str) -> list[str]:
    names = []
    if team == "Todas":
        names = [member["Nome"] for member in staff if member.get("Ativo", True)]
    else:
        names = [
            member["Nome"]
            for member in staff
            if member.get("Ativo", True) and member.get("Equipa", "").strip() == team
        ]
    vacation_names = sorted({vacation["employee_name"] for vacation in vacations})
    ordered = []
    seen = set()
    for name in names + vacation_names:
        if name not in seen:
            seen.add(name)
            ordered.append(name)
    return ordered


def approved_days_for_employee(vacations: list[dict], employee_name: str, year: int, month: int) -> set[int]:
    days: set[int] = set()
    month_start = date(year, month, 1)
    month_end = date(year, month, monthrange(year, month)[1])
    for vacation in vacations:
        if vacation["employee_name"] != employee_name:
            continue
        start = max(vacation["start_date"], month_start)
        end = min(vacation["end_date"], month_end)
        current = start
        while current <= end:
            days.add(current.day)
            current += timedelta(days=1)
    return days


def build_margarida_workbook(
    vacations: list[dict],
    staff: list[dict],
    *,
    year: int,
    team: str,
    title_name: str = "Margarida",
) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError("A dependência openpyxl não está disponível.") from exc

    filtered_vacations = vacations_for_year(vacations, year, team)
    employees = employee_names_for_map(staff, filtered_vacations, team)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = str(year)
    sheet.freeze_panes = "B8"

    title_fill = PatternFill("solid", fgColor="D9E8FB")
    header_fill = PatternFill("solid", fgColor="EAF2FB")
    x_fill = PatternFill("solid", fgColor="E07A7A")
    month_fill = PatternFill("solid", fgColor="D7E6F7")
    thin = Side(style="thin", color="B7C6D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    max_days = 31
    last_day_col = 1 + max_days
    final_merge_col = last_day_col

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=final_merge_col)
    top_title = sheet.cell(row=2, column=1, value=f"Calendário {year}")
    top_title.font = Font(size=14, bold=True)
    top_title.alignment = center
    top_title.fill = title_fill
    top_title.border = border

    sheet.merge_cells(start_row=3, start_column=8, end_row=3, end_column=21)
    map_title = sheet.cell(row=3, column=8, value=f"{title_name} - MAPA FÉRIAS {team.upper() if team != 'Todas' else 'EQUIPA'}")
    map_title.font = Font(size=12, bold=True)
    map_title.alignment = center
    map_title.fill = header_fill
    map_title.border = border

    sheet.column_dimensions["A"].width = 24
    for column in range(2, last_day_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 4

    row = 5
    for month in range(1, 13):
        days_in_month = monthrange(year, month)[1]
        first_day = date(year, month, 1)

        weekday_row = row
        month_row = row + 1
        employee_start_row = row + 2

        sheet.cell(row=weekday_row, column=1, value="")
        sheet.cell(row=month_row, column=1, value=month_name_pt(month))
        sheet.cell(row=month_row, column=1).fill = month_fill
        sheet.cell(row=month_row, column=1).font = Font(bold=True)

        for day in range(1, max_days + 1):
            weekday_cell = sheet.cell(row=weekday_row, column=day + 1)
            number_cell = sheet.cell(row=month_row, column=day + 1)
            weekday_cell.alignment = center
            number_cell.alignment = center
            weekday_cell.border = border
            number_cell.border = border
            if day <= days_in_month:
                current = date(year, month, day)
                weekday_cell.value = PT_WEEKDAYS[current.weekday()]
                number_cell.value = day
                if current.weekday() >= 5:
                    weekday_cell.fill = header_fill
                    number_cell.fill = header_fill
            else:
                weekday_cell.value = ""
                number_cell.value = ""

        for index, employee in enumerate(employees):
            current_row = employee_start_row + index
            name_cell = sheet.cell(row=current_row, column=1, value=employee)
            name_cell.border = border
            approved_days = approved_days_for_employee(filtered_vacations, employee, year, month)
            for day in range(1, max_days + 1):
                cell = sheet.cell(row=current_row, column=day + 1)
                cell.alignment = center
                cell.border = border
                if day <= days_in_month and day in approved_days:
                    cell.value = "X"
                    cell.fill = x_fill
                    cell.font = Font(bold=True, color="FFFFFF")

        row = employee_start_row + len(employees) + 1

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
